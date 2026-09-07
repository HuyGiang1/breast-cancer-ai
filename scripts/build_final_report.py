#!/usr/bin/env python3

from __future__ import annotations

import argparse
import re
import tempfile
import zipfile
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path

from PIL import Image as PILImage
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY, TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (
    Image,
    KeepTogether,
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_SOURCE = ROOT / "docs/report/FINAL_RESEARCH_REPORT_VI.md"
OUTPUT_STEM = "NGHIEN_CUU_CAC_MO_HINH_NHAN_DANG_PHAN_LOAI_KHOI_U_VU_AC_TINH"


@dataclass
class Block:
    kind: str
    value: object = None
    level: int = 0


def clean_inline(value: str) -> str:
    value = re.sub(r"`([^`]+)`", r"\1", value)
    value = re.sub(r"\*\*([^*]+)\*\*", r"\1", value)
    return value.strip()


def parse_markdown(source: Path) -> list[Block]:
    lines = source.read_text(encoding="utf-8").splitlines()
    blocks: list[Block] = []
    index = 0
    while index < len(lines):
        raw = lines[index]
        text = raw.strip()
        if not text:
            index += 1
            continue
        if text == "<!-- PAGEBREAK -->":
            blocks.append(Block("pagebreak"))
            index += 1
            continue
        heading = re.match(r"^(#{1,4})\s+(.+)$", text)
        if heading:
            blocks.append(Block("heading", clean_inline(heading.group(2)), len(heading.group(1))))
            index += 1
            continue
        image = re.match(r"^!\[(.*?)\]\((.*?)\)$", text)
        if image:
            image_path = (source.parent / image.group(2)).resolve()
            blocks.append(Block("image", (image_path, clean_inline(image.group(1)))))
            index += 1
            continue
        if text.startswith("|"):
            table_lines: list[str] = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                table_lines.append(lines[index].strip())
                index += 1
            rows = [[clean_inline(cell) for cell in row.strip("|").split("|")] for row in table_lines]
            if len(rows) > 1 and all(re.fullmatch(r":?-{3,}:?", cell) for cell in rows[1]):
                rows.pop(1)
            blocks.append(Block("table", rows))
            continue
        if re.match(r"^[-*]\s+", text):
            items: list[str] = []
            while index < len(lines) and re.match(r"^[-*]\s+", lines[index].strip()):
                items.append(clean_inline(re.sub(r"^[-*]\s+", "", lines[index].strip())))
                index += 1
            blocks.append(Block("bullets", items))
            continue
        if re.match(r"^\d+\.\s+", text):
            items: list[str] = []
            while index < len(lines) and re.match(r"^\d+\.\s+", lines[index].strip()):
                items.append(clean_inline(re.sub(r"^\d+\.\s+", "", lines[index].strip())))
                index += 1
            blocks.append(Block("numbers", items))
            continue
        paragraph = [text]
        index += 1
        while index < len(lines):
            nxt = lines[index].strip()
            if not nxt:
                break
            if nxt == "<!-- PAGEBREAK -->" or nxt.startswith("#") or nxt.startswith("|"):
                break
            if nxt.startswith("![") or re.match(r"^[-*]\s+", nxt) or re.match(r"^\d+\.\s+", nxt):
                break
            paragraph.append(nxt)
            index += 1
        blocks.append(Block("paragraph", clean_inline(" ".join(paragraph))))
    return blocks


def set_cell_shading(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, end])


def configure_docx(document: Document) -> None:
    section = document.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(3.0)
    section.right_margin = Cm(2.0)
    add_page_number(section.footer.paragraphs[0])

    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.font.size = Pt(12.5)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.line_spacing = 1.35
    normal.paragraph_format.space_after = Pt(7)

    for level, size, color in ((1, 18, "123047"), (2, 15, "123047"), (3, 13, "176C62")):
        style = document.styles[f"Heading {level}"]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = __import__("docx").shared.RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(12)
        style.paragraph_format.space_after = Pt(8)
        style.paragraph_format.keep_with_next = True

    caption = document.styles.add_style("Figure Caption", WD_STYLE_TYPE.PARAGRAPH)
    caption.font.name = "Times New Roman"
    caption.font.size = Pt(10.5)
    caption.font.italic = True
    caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.space_after = Pt(9)


def add_docx_image(document: Document, path: Path, caption: str) -> None:
    if not path.exists():
        raise FileNotFoundError(path)
    with PILImage.open(path) as image:
        width, height = image.size
    max_width, max_height = 15.5, 17.5
    scale = min(max_width / width, max_height / height)
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run().add_picture(str(path), width=Cm(width * scale), height=Cm(height * scale))
    document.add_paragraph(caption, style="Figure Caption")


def normalize_docx_archive(path: Path) -> None:
    fixed_time = (2026, 9, 7, 0, 0, 0)
    with tempfile.NamedTemporaryFile(dir=path.parent, suffix=".docx", delete=False) as handle:
        normalized_path = Path(handle.name)
    try:
        with zipfile.ZipFile(path, "r") as source, zipfile.ZipFile(
            normalized_path, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=9
        ) as target:
            for source_info in sorted(source.infolist(), key=lambda item: item.filename):
                info = zipfile.ZipInfo(source_info.filename, fixed_time)
                info.compress_type = zipfile.ZIP_DEFLATED
                info.external_attr = source_info.external_attr
                info.create_system = source_info.create_system
                target.writestr(info, source.read(source_info.filename))
        normalized_path.replace(path)
    finally:
        normalized_path.unlink(missing_ok=True)


def build_docx(blocks: list[Block], output: Path) -> None:
    document = Document()
    configure_docx(document)
    title_page = True
    for block in blocks:
        if block.kind == "pagebreak":
            document.add_page_break()
            title_page = False
        elif block.kind == "heading":
            paragraph = document.add_heading(str(block.value), level=min(block.level, 3))
            if title_page:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                paragraph.paragraph_format.space_before = Pt(28 if block.level == 1 else 8)
        elif block.kind == "paragraph":
            paragraph = document.add_paragraph(str(block.value))
            if title_page:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                paragraph.paragraph_format.space_before = Pt(10)
        elif block.kind in {"bullets", "numbers"}:
            style = "List Bullet" if block.kind == "bullets" else "List Number"
            for item in block.value:
                paragraph = document.add_paragraph(str(item), style=style)
                paragraph.paragraph_format.space_after = Pt(3)
        elif block.kind == "table":
            rows = block.value
            table = document.add_table(rows=len(rows), cols=len(rows[0]))
            table.style = "Table Grid"
            for row_index, row in enumerate(rows):
                for col_index, value in enumerate(row):
                    cell = table.cell(row_index, col_index)
                    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    cell.text = value
                    for paragraph in cell.paragraphs:
                        paragraph.paragraph_format.space_after = Pt(2)
                        paragraph.paragraph_format.line_spacing = 1.0
                        for run in paragraph.runs:
                            run.font.name = "Times New Roman"
                            run.font.size = Pt(9.5)
                            run.font.bold = row_index == 0
                    if row_index == 0:
                        set_cell_shading(cell, "DCEDEA")
                row_element = table.rows[row_index]._tr
                properties = row_element.get_or_add_trPr()
                cant_split = OxmlElement("w:cantSplit")
                properties.append(cant_split)
                table.rows[row_index].height_rule = WD_ROW_HEIGHT_RULE.AUTO
            table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
            document.add_paragraph()
        elif block.kind == "image":
            path, caption = block.value
            add_docx_image(document, path, caption)

    properties = document.core_properties
    properties.title = "Nghiên cứu các mô hình nhận dạng, phân loại các khối u vú ác tính"
    properties.subject = "Báo cáo nghiên cứu khoa học sinh viên"
    properties.author = "Nguyễn Bá Duy; Trần Mỹ Anh; Hoàng Nhật Anh; Nguyễn Huy Giang; Ngô Tiến Đạt"
    fixed_date = datetime(2026, 9, 7, tzinfo=timezone.utc)
    properties.created = fixed_date
    properties.modified = fixed_date
    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)
    normalize_docx_archive(output)


def register_pdf_fonts() -> None:
    font_dir = Path("/System/Library/Fonts/Supplemental")
    pdfmetrics.registerFont(TTFont("TNR", font_dir / "Times New Roman.ttf"))
    pdfmetrics.registerFont(TTFont("TNR-Bold", font_dir / "Times New Roman Bold.ttf"))
    pdfmetrics.registerFont(TTFont("TNR-Italic", font_dir / "Times New Roman Italic.ttf"))
    pdfmetrics.registerFontFamily("TNR", normal="TNR", bold="TNR-Bold", italic="TNR-Italic")


def pdf_styles():
    styles = getSampleStyleSheet()
    return {
        "body": ParagraphStyle("Body", fontName="TNR", fontSize=11.5, leading=17, alignment=TA_JUSTIFY, spaceAfter=7),
        "h1": ParagraphStyle("H1", fontName="TNR-Bold", fontSize=17, leading=21, textColor=colors.HexColor("#123047"), spaceBefore=10, spaceAfter=10),
        "h2": ParagraphStyle("H2", fontName="TNR-Bold", fontSize=14, leading=18, textColor=colors.HexColor("#123047"), spaceBefore=9, spaceAfter=7),
        "h3": ParagraphStyle("H3", fontName="TNR-Bold", fontSize=12, leading=16, textColor=colors.HexColor("#176C62"), spaceBefore=7, spaceAfter=5),
        "title": ParagraphStyle("Title", fontName="TNR-Bold", fontSize=21, leading=28, alignment=TA_CENTER, textColor=colors.HexColor("#123047"), spaceBefore=45, spaceAfter=22),
        "cover": ParagraphStyle("Cover", fontName="TNR", fontSize=12.5, leading=20, alignment=TA_CENTER, spaceAfter=12),
        "caption": ParagraphStyle("Caption", fontName="TNR-Italic", fontSize=9.5, leading=12, alignment=TA_CENTER, spaceAfter=10),
        "list": ParagraphStyle("List", fontName="TNR", fontSize=11.5, leading=16, leftIndent=15, firstLineIndent=-8, spaceAfter=4),
        "cell": ParagraphStyle("Cell", fontName="TNR", fontSize=8.7, leading=10.5, alignment=TA_LEFT),
        "cell_header": ParagraphStyle("CellHeader", fontName="TNR-Bold", fontSize=8.7, leading=10.5, alignment=TA_CENTER),
    }


def add_pdf_footer(canvas, document) -> None:
    canvas.saveState()
    canvas.setFont("TNR", 9)
    canvas.setFillColor(colors.HexColor("#52636D"))
    canvas.drawCentredString(A4[0] / 2, 1.15 * cm, str(document.page))
    canvas.restoreState()


def build_pdf(blocks: list[Block], output: Path) -> None:
    register_pdf_fonts()
    styles = pdf_styles()
    story = []
    title_page = True
    for block in blocks:
        if block.kind == "pagebreak":
            story.append(PageBreak())
            title_page = False
        elif block.kind == "heading":
            style = styles["title"] if title_page and block.level == 1 else styles[f"h{min(block.level, 3)}"]
            story.append(Paragraph(str(block.value), style))
        elif block.kind == "paragraph":
            story.append(Paragraph(str(block.value), styles["cover"] if title_page else styles["body"]))
        elif block.kind in {"bullets", "numbers"}:
            for item_index, item in enumerate(block.value, start=1):
                marker = "•" if block.kind == "bullets" else f"{item_index}."
                story.append(Paragraph(f"{marker} {item}", styles["list"]))
        elif block.kind == "table":
            rows = []
            for row_index, row in enumerate(block.value):
                style = styles["cell_header"] if row_index == 0 else styles["cell"]
                rows.append([Paragraph(cell, style) for cell in row])
            available = A4[0] - 5 * cm
            table = Table(rows, repeatRows=1, colWidths=[available / len(rows[0])] * len(rows[0]), hAlign="CENTER")
            table.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#DCEDEA")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.HexColor("#123047")),
                ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#9AA8AE")),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("LEFTPADDING", (0, 0), (-1, -1), 5),
                ("RIGHTPADDING", (0, 0), (-1, -1), 5),
                ("TOPPADDING", (0, 0), (-1, -1), 4),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ]))
            story.extend([table, Spacer(1, 8)])
        elif block.kind == "image":
            path, caption = block.value
            if not path.exists():
                raise FileNotFoundError(path)
            with PILImage.open(path) as image:
                width, height = image.size
            max_width, max_height = 15.5 * cm, 15.5 * cm
            scale = min(max_width / width, max_height / height)
            figure = Image(str(path), width=width * scale, height=height * scale)
            figure.hAlign = "CENTER"
            story.append(KeepTogether([figure, Spacer(1, 4), Paragraph(caption, styles["caption"])]))

    output.parent.mkdir(parents=True, exist_ok=True)
    document = SimpleDocTemplate(
        str(output),
        pagesize=A4,
        leftMargin=3 * cm,
        rightMargin=2 * cm,
        topMargin=2.2 * cm,
        bottomMargin=2 * cm,
        title="Nghiên cứu các mô hình nhận dạng, phân loại các khối u vú ác tính",
        author="Nguyễn Bá Duy; Trần Mỹ Anh; Hoàng Nhật Anh; Nguyễn Huy Giang; Ngô Tiến Đạt",
        subject="Báo cáo nghiên cứu khoa học sinh viên",
        invariant=1,
    )
    document.build(story, onFirstPage=add_pdf_footer, onLaterPages=add_pdf_footer)


def main() -> None:
    parser = argparse.ArgumentParser(description="Build the official final report DOCX and PDF.")
    parser.add_argument("--source", type=Path, default=DEFAULT_SOURCE)
    parser.add_argument("--output-dir", type=Path, default=ROOT / "docs/report")
    args = parser.parse_args()
    blocks = parse_markdown(args.source.resolve())
    build_docx(blocks, args.output_dir / f"{OUTPUT_STEM}.docx")
    build_pdf(blocks, args.output_dir / f"{OUTPUT_STEM}.pdf")
    print(f"Built final report from {args.source}")


if __name__ == "__main__":
    main()
