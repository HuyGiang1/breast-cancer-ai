#!/usr/bin/env python3

from __future__ import annotations

import json
import re
import sys
from pathlib import Path

from docx import Document
from pypdf import PdfReader


ROOT = Path(__file__).resolve().parents[1]
REPORT_DIR = ROOT / "docs/report"
STEM = "NGHIEN_CUU_CAC_MO_HINH_NHAN_DANG_PHAN_LOAI_KHOI_U_VU_AC_TINH"


def normalized(text: str) -> str:
    return re.sub(r"\s+", " ", text).strip().lower()


def docx_text(document: Document) -> str:
    content = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            content.extend(cell.text for cell in row.cells)
    return "\n".join(content)


def require(condition: bool, message: str, failures: list[str]) -> None:
    if not condition:
        failures.append(message)


def main() -> int:
    source_path = REPORT_DIR / "FINAL_RESEARCH_REPORT_VI.md"
    docx_path = REPORT_DIR / f"{STEM}.docx"
    pdf_path = REPORT_DIR / f"{STEM}.pdf"
    snapshot_path = ROOT / "experiments/final/FINAL_RESULTS_SNAPSHOT.json"
    public_paths = [
        ROOT / "README.md",
        source_path,
        ROOT / "docs/RELEASE_NOTES.md",
        ROOT / "docs/DEPLOYMENT.md",
        ROOT / "docs/DEPLOYMENT_RUNBOOK.md",
        ROOT / "docs/PROJECT_PROGRESS.md",
        ROOT / "docs/PROJECT_STATUS.md",
        ROOT / "docs/AGENT_HANDOFF.md",
    ]

    failures: list[str] = []
    for path in [source_path, docx_path, pdf_path, snapshot_path]:
        require(path.is_file() and path.stat().st_size > 0, f"Missing or empty: {path}", failures)
    if failures:
        print("FINAL REPORT VALIDATION: FAIL")
        print("\n".join(f"- {failure}" for failure in failures))
        return 1

    snapshot = json.loads(snapshot_path.read_text(encoding="utf-8"))
    source = source_path.read_text(encoding="utf-8")
    document = Document(docx_path)
    docx = docx_text(document)
    reader = PdfReader(pdf_path)
    pdf_pages = [page.extract_text() or "" for page in reader.pages]
    pdf = "\n".join(pdf_pages)

    report_texts = {
        "source": normalized(source),
        "DOCX": normalized(docx),
        "PDF": normalized(pdf),
    }
    required_facts = [
        "nghiên cứu các mô hình nhận dạng, phân loại các khối u vú ác tính",
        "nguyễn bá duy",
        "trần mỹ anh",
        "hoàng nhật anh",
        "nguyễn huy giang",
        "ngô tiến đạt",
        "đoàn thị thanh hằng",
        "logistic regression",
        "efficientnet-b0",
        "raw_probability >= 0.515",
        "clinical_use=false",
        "experimental_only",
    ]
    for output_name, text in report_texts.items():
        for fact in required_facts:
            require(fact in text, f"{output_name} missing fact: {fact}", failures)

    require(snapshot["dataset"]["wdbc"]["total_samples"] == 569, "Snapshot WDBC total changed", failures)
    require(snapshot["dataset"]["cbis_ddsm"]["manifest_rows"] == 5118, "Snapshot manifest rows changed", failures)
    require(snapshot["ml"]["primary_candidate"] == "Logistic Regression", "Snapshot ML candidate changed", failures)
    require(snapshot["dl"]["retained_candidate"].startswith("EfficientNet-B0"), "Snapshot DL candidate changed", failures)
    require(snapshot["roi"]["decision"] == "ROI-C", "Snapshot ROI decision changed", failures)
    require(float(snapshot["threshold"]["ml"]["Logistic Regression"]) == 0.36, "Snapshot ML threshold changed", failures)
    require(abs(float(snapshot["threshold"]["dl"]) - 0.515) < 1e-9, "Snapshot DL threshold changed", failures)

    require(25 <= len(reader.pages) <= 60, f"Unexpected PDF page count: {len(reader.pages)}", failures)
    for page_number, page_text in enumerate(pdf_pages, start=1):
        require(len(normalized(page_text)) >= 80, f"PDF page {page_number} appears blank", failures)
    require(len(document.inline_shapes) >= 15, "DOCX is missing final figures/screenshots", failures)
    require(len(document.tables) >= 7, "DOCX is missing final tables", failures)

    forbidden = {
        "calibrated threshold contract": r"calibrated(?:_|\s+)probability\s*>=\s*0[.,]515",
        "5118 images claim": r"5[.,]?118\s+(?:independent\s+)?images",
        "5118 patients claim": r"5[.,]?118\s+(?:independent\s+)?patients",
        "clinical-grade claim": r"clinical[- ]grade",
        "validated multimodal claim": r"(?<!no )(?<!not )validated multimodal",
        "React architecture claim": r"\breact\b",
        "Next.js architecture claim": r"\bnext\.js\b",
        "missing-server wording": r"vps (?:unavailable|not provisioned)|no server exists",
    }
    combined_public = "\n".join(path.read_text(encoding="utf-8") for path in public_paths if path.exists())
    for label, pattern in forbidden.items():
        require(not re.search(pattern, combined_public, re.IGNORECASE), f"Forbidden final-facing wording: {label}", failures)

    if failures:
        print("FINAL REPORT VALIDATION: FAIL")
        print("\n".join(f"- {failure}" for failure in failures))
        return 1

    print("FINAL REPORT VALIDATION: PASS")
    print(f"PDF pages: {len(reader.pages)}")
    print(f"DOCX figures: {len(document.inline_shapes)}")
    print(f"DOCX tables: {len(document.tables)}")
    print("Scientific contract: snapshot/source/DOCX/PDF consistent")
    return 0


if __name__ == "__main__":
    sys.exit(main())
